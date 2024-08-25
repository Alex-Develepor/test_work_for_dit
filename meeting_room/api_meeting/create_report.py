from docx import Document

from .models import LogForMeetingRoom


def create_report():
    document = Document()
    objects = LogForMeetingRoom.objects.all()
    table = document.add_table(len(objects), 4)
    for i, obj in enumerate(objects):
        table.cell(i, 0).text = obj.user.username
        table.cell(i, 1).text = obj.period_reservation
        table.cell(i, 2).text = obj.informations
        table.cell(i, 3).text = obj.meetingroom.name
    document.save('report.docx')

